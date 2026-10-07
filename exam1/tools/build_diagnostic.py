"""Build accessible Word copies from the two diagnostic Markdown sources.

Run from any directory with Python 3 and python-docx installed.
The Markdown files are the source of truth. PDF export is a separate step.
"""
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]


def add_table(doc, rows):
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = {3: [1.15, 1.45, 3.9],
              4: [1.25, 1.75, 1.75, 1.75],
              5: [1.2, 1.1, 1.1, 1.1, 2.0]}[len(rows[0])]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + side)
        for key, val in [('val', 'single'), ('sz', '6'), ('color', 'D9D9D9')]:
            edge.set(qn('w:' + key), val)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for col, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col == 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.bold = index == 0
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E8E8E8' if index == 0 else 'FFFFFF')
            props.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                item = OxmlElement('w:' + side)
                item.set(qn('w:w'), '100')
                item.set(qn('w:type'), 'dxa')
                margins.append(item)
            props.append(margins)


def inline(paragraph, text):
    for piece in re.split(r'(\[[^\]]+\]\([^)]+\))', text):
        match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', piece)
        if not match:
            paragraph.add_run(piece.replace('`', ''))
            continue
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), paragraph.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        underline = OxmlElement('w:u')
        underline.set(qn('w:val'), 'single')
        props.append(underline)
        run.append(props)
        value = OxmlElement('w:t')
        value.text = match[1]
        run.append(value)
        link.append(run)
        paragraph._p.append(link)


def build(stem):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.55 if stem == 'questions' else 0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(11 if stem == 'questions' else 12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(3 if stem == 'questions' else 5)
    normal.paragraph_format.line_spacing = 1.0
    for name, size in [('Title', 22), ('Heading 1', 16), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(11.5 if stem == 'questions' and name == 'Heading 2' else size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt((6 if stem == 'questions' else 8) if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(3 if stem == 'questions' else 5)
        style.paragraph_format.keep_with_next = True
    lang = OxmlElement('w:lang')
    lang.set(qn('w:val'), 'en-US')
    normal.element.get_or_add_rPr().append(lang)
    doc.core_properties.title = ('Exam 1 foundation check' if stem == 'questions'
                                 else 'Exam 1 foundation check answer guide')
    doc.core_properties.subject = 'CSC225 student practice and review routing'
    doc.core_properties.author = 'CSC225 TA resources'
    footer = section.footer.paragraphs[0]
    footer.add_run('CSC225  |  ' + ('Foundation check' if stem == 'questions' else 'Answer guide') + '  |  Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)
    lines = (ROOT / 'diagnostic' / (stem + '.md')).read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        i += 1
        if not line:
            continue
        if line == '<!-- pagebreak -->':
            doc.add_page_break()
        elif line.startswith('```'):
            code = []
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.keep_together = True
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(3 if stem == 'questions' else 5)
            run = p.add_run('\n'.join(code))
            run.font.name = 'Consolas'
            run.font.size = Pt(10.5 if stem == 'questions' else 11)
        elif line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            doc.add_paragraph(line[3:], 'Heading 1')
        elif line.startswith('### '):
            doc.add_paragraph(line[4:], 'Heading 2')
        elif line.startswith('- '):
            inline(doc.add_paragraph(style='List Bullet'), line[2:])
        elif line.startswith('|'):
            raw_rows = [line]
            while i < len(lines) and lines[i].startswith('|'):
                raw_rows.append(lines[i])
                i += 1
            rows = [[cell.strip() for cell in row.strip('|').split('|')]
                    for row in raw_rows if not re.fullmatch(r'[| :\-]+', row)]
            add_table(doc, rows)
        else:
            block = [line]
            while i < len(lines) and lines[i] and not lines[i].startswith(('#', '```', '<!--', '- ', '|')):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            inline(p, '\n'.join(block))
            p.paragraph_format.keep_together = True
            if stem == 'questions' and not line.startswith(('Answer:', 'Reason', 'Output:', 'Values of', 'Next:', '____')):
                p.paragraph_format.keep_with_next = True
    target = ROOT / 'diagnostic' / (stem + '.docx')
    doc.save(target)
    print(target)


if __name__ == '__main__':
    build('questions')
    build('answers')
